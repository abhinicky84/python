from __future__ import annotations

import base64
from dataclasses import dataclass
from io import BytesIO
import re
from xml.sax.saxutils import escape

from app.schemas.architecture import ArchitectureAnalysis


@dataclass(frozen=True)
class DiagramNode:
    key: str
    label: str
    category: str


@dataclass(frozen=True)
class DiagramEdge:
    source: str
    target: str
    label: str
    style: str = "solid"


class DiagramGenerator:
    CATEGORY_ORDER = [
        "channels",
        "experience",
        "identity",
        "integration",
        "applications",
        "systems",
        "data",
        "observability",
    ]

    CATEGORY_LABELS = {
        "channels": "User Channels",
        "experience": "Experience Layer",
        "identity": "Security & Identity",
        "integration": "Integration Layer",
        "applications": "Application Services",
        "systems": "Enterprise Platforms",
        "data": "Data & Analytics",
        "observability": "Observability",
    }

    CATEGORY_STYLES = {
        "channels": "fill:#fef3c7,stroke:#d97706,color:#111827",
        "experience": "fill:#dbeafe,stroke:#2563eb,color:#111827",
        "identity": "fill:#ede9fe,stroke:#7c3aed,color:#111827",
        "integration": "fill:#dcfce7,stroke:#16a34a,color:#111827",
        "applications": "fill:#fae8ff,stroke:#c026d3,color:#111827",
        "systems": "fill:#fee2e2,stroke:#dc2626,color:#111827",
        "data": "fill:#cffafe,stroke:#0891b2,color:#111827",
        "observability": "fill:#e5e7eb,stroke:#4b5563,color:#111827",
    }

    DRAWIO_LANE_STYLES = {
        "channels": "swimlane;html=1;rounded=1;fillColor=#fef3c7;strokeColor=#d97706;fontStyle=1;",
        "experience": "swimlane;html=1;rounded=1;fillColor=#dbeafe;strokeColor=#2563eb;fontStyle=1;",
        "identity": "swimlane;html=1;rounded=1;fillColor=#ede9fe;strokeColor=#7c3aed;fontStyle=1;",
        "integration": "swimlane;html=1;rounded=1;fillColor=#dcfce7;strokeColor=#16a34a;fontStyle=1;",
        "applications": "swimlane;html=1;rounded=1;fillColor=#fae8ff;strokeColor=#c026d3;fontStyle=1;",
        "systems": "swimlane;html=1;rounded=1;fillColor=#fee2e2;strokeColor=#dc2626;fontStyle=1;",
        "data": "swimlane;html=1;rounded=1;fillColor=#cffafe;strokeColor=#0891b2;fontStyle=1;",
        "observability": "swimlane;html=1;rounded=1;fillColor=#e5e7eb;strokeColor=#4b5563;fontStyle=1;",
    }

    DRAWIO_NODE_STYLES = {
        "channels": "rounded=1;whiteSpace=wrap;html=1;fillColor=#fff7d6;strokeColor=#d97706;shadow=1;spacing=10;fontStyle=1;",
        "experience": "rounded=1;whiteSpace=wrap;html=1;fillColor=#edf5ff;strokeColor=#2563eb;shadow=1;spacing=10;fontStyle=1;",
        "identity": "rounded=1;whiteSpace=wrap;html=1;fillColor=#f5f3ff;strokeColor=#7c3aed;shadow=1;spacing=10;fontStyle=1;",
        "integration": "rounded=1;whiteSpace=wrap;html=1;fillColor=#eafcf0;strokeColor=#16a34a;fontStyle=1;shadow=1;spacing=10;",
        "applications": "rounded=1;whiteSpace=wrap;html=1;fillColor=#fdf0ff;strokeColor=#c026d3;shadow=1;spacing=10;fontStyle=1;",
        "systems": "rounded=1;whiteSpace=wrap;html=1;fillColor=#fff1f2;strokeColor=#dc2626;shadow=1;spacing=10;fontStyle=1;",
        "data": "rounded=1;whiteSpace=wrap;html=1;fillColor=#ecfeff;strokeColor=#0891b2;shadow=1;spacing=10;fontStyle=1;",
        "observability": "rounded=1;whiteSpace=wrap;html=1;fillColor=#f3f4f6;strokeColor=#4b5563;shadow=1;spacing=10;fontStyle=1;",
    }

    DRAWIO_NODE_DETAILS = {
        "users": ["Customers", "Employees", "B2B teams"],
        "partners": ["Suppliers", "Agencies", "Distributors"],
        "web": ["Web", "Mobile", "Portal entry points"],
        "frontend": ["Presentation APIs", "BFF", "Session handling"],
        "aem": ["Content fragments", "Experience fragments", "Authoring", "Workflows", "DAM integration", "Publishing"],
        "aem_assets": ["Asset repository", "Metadata", "Renditions", "Governance", "Search", "Distribution"],
        "portal": ["Content orchestration", "Digital touchpoints"],
        "commerce": ["Catalog", "Pricing", "Cart", "Checkout", "Promotions", "Inventory"],
        "edge": ["Caching", "Acceleration", "Global routing"],
        "personalization": ["Segmentation", "Decisioning", "Recommendations", "Targeting", "A/B testing", "Next best action"],
        "target": ["Offers", "Targeting", "Experiments", "Personalized content"],
        "dynamic_media": ["Image delivery", "Video", "Renditions", "Media optimization"],
        "seo": ["Structured metadata", "Sitemaps", "Discoverability"],
        "entra": ["SSO", "OAuth / OIDC", "Role-based access"],
        "network": ["Private Link", "VNet isolation", "Ingress controls"],
        "edge_api": ["WAF-aligned rules", "Rate limiting", "Edge validation"],
        "api_gateway": ["North-south APIs", "Contract mediation", "Traffic shaping"],
        "apim": ["Policies", "Developer access", "Security enforcement"],
        "orchestration": ["Process flows", "Transformation", "System coordination"],
        "service_bus": ["Queues", "Topics", "Reliable delivery"],
        "event_grid": ["Event distribution", "Webhook fan-out", "Notifications"],
        "domain_services": ["Core business logic", "Composable domains", "Reusable APIs"],
        "microservices": ["Bounded contexts", "Independent release", "Domain autonomy"],
        "container_apps": ["Scalable runtime", "Revision rollout", "Service endpoints"],
        "aks": ["Kubernetes runtime", "Platform controls", "Advanced workloads"],
        "shared_services": ["Common capabilities", "Enterprise reuse", "Shared policies"],
        "service_mesh": ["mTLS", "East-west controls", "Traffic governance"],
        "ai_foundry": ["Prompt flows", "Model access", "AI orchestration"],
        "sap": ["Order management", "Product master", "Pricing", "Inventory", "Finance", "Fulfillment"],
        "sap_cpi": ["Integration flows", "Mappings", "Adapters", "B2B/EDI"],
        "sap_btp": ["Extension apps", "Workflows", "Event mesh", "Data services"],
        "fiori": ["UX apps", "Launchpad", "Role-based access", "Approvals"],
        "salesforce": ["CRM", "Service workflows", "Customer engagement"],
        "martech": ["Campaigns", "Journeys", "Audience activation"],
        "loyalty": ["Rewards", "Membership", "Offers"],
        "search_platform": ["Site search", "Faceting", "Relevance tuning"],
        "data_platform": ["Lakehouse", "BI", "Data products"],
        "lob": ["Back-office apps", "Operational systems", "Legacy integration"],
        "operational_data": ["Transactional store", "Low-latency access", "Operational views"],
        "canonical_data": ["Shared business schema", "Event contracts", "Master data exchange"],
        "customer_profile": ["Unified profile", "Segmentation", "Activation context"],
        "search_index": ["Indexing", "Search feeds", "SEO publishing"],
        "analytics": ["Reporting", "Insights", "Decision support"],
        "adobe_analytics": ["Traffic analytics", "Conversion", "Customer journeys", "Dashboards"],
        "aep": ["Profile unification", "Schemas", "Datasets", "Activation"],
        "rt_cdp": ["Audience building", "Real-time profiles", "Activation", "Consent"],
        "azure_cloud": ["App services", "Integration", "Data platform", "Identity"],
        "aws_cloud": ["Storage", "CDN", "Compute", "Messaging"],
        "gcp_cloud": ["Analytics", "AI/ML", "Storage", "Data processing"],
        "monitor": ["Metrics", "Tracing", "Alerting"],
        "governance": ["Audit trail", "Compliance", "Operational guardrails"],
    }

    DRAWIO_PRIMARY_EDGE_LABELS = {
        ("users", "web"),
        ("web", "frontend"),
        ("frontend", "apim"),
        ("entra", "apim"),
        ("apim", "orchestration"),
        ("orchestration", "domain_services"),
        ("service_bus", "microservices"),
        ("domain_services", "canonical_data"),
        ("canonical_data", "analytics"),
        ("analytics", "personalization"),
        ("personalization", "customer_profile"),
        ("canonical_data", "search_index"),
    }

    DRAWIO_LAYOUT_HINTS = {
        "users": (0, 0),
        "partners": (0, 1),
        "web": (1, 0),
        "edge": (1, 0),
        "frontend": (2, 0),
        "entra": (2, 0),
        "network": (1, 0),
        "aem": (3, 0),
        "portal": (3, 0),
        "commerce": (4, 0),
        "seo": (4, 1),
        "personalization": (5, 0),
        "edge_api": (1, 0),
        "api_gateway": (2, 0),
        "apim": (3, 0),
        "orchestration": (4, 0),
        "service_bus": (5, 0),
        "event_grid": (5, 1),
        "domain_services": (4, 0),
        "container_apps": (4, 1),
        "microservices": (5, 0),
        "aks": (5, 1),
        "shared_services": (6, 0),
        "service_mesh": (6, 1),
        "ai_foundry": (6, 0),
        "sap": (4, 0),
        "lob": (4, 1),
        "salesforce": (5, 0),
        "martech": (6, 0),
        "loyalty": (6, 1),
        "search_platform": (7, 0),
        "data_platform": (7, 1),
        "operational_data": (4, 0),
        "canonical_data": (5, 0),
        "customer_profile": (6, 0),
        "search_index": (6, 1),
        "analytics": (7, 0),
        "monitor": (5, 0),
        "governance": (6, 0),
    }

    PPTX_NODE_FILL = {
        "channels": "F5F7FB",
        "experience": "FFF2CC",
        "identity": "F4E9FF",
        "integration": "DAE8FC",
        "applications": "D5E8D4",
        "systems": "F8CECC",
        "data": "E1D5E7",
        "observability": "E7EDF4",
    }

    PPTX_NODE_LINE = {
        "channels": "64748B",
        "experience": "D6B656",
        "identity": "8B5CF6",
        "integration": "6C8EBF",
        "applications": "82B366",
        "systems": "B85450",
        "data": "9673A6",
        "observability": "64748B",
    }

    def generate(self, user_input: str, model_response: str, analysis: ArchitectureAnalysis) -> tuple[str, str, str]:
        nodes = self._build_nodes(user_input, model_response, analysis)
        edges = self._build_edges(nodes)
        return (
            self._build_drawio_xml(nodes, edges),
            self._build_pptx_base64(nodes, edges),
            self._build_pptx_preview_svg(nodes, edges),
        )

    def _build_nodes(
        self,
        user_input: str,
        model_response: str,
        analysis: ArchitectureAnalysis,
    ) -> list[DiagramNode]:
        text = f"{user_input}\n{model_response}"
        nodes: list[DiagramNode] = []
        cloud_provider = analysis.cloud_provider.casefold()

        self._add_node(nodes, "users", "Business Users", "channels")
        self._add_node(nodes, "web", "Web / Mobile Channels", "channels")
        if self._mentions_any(text, ["partner", "dealer", "vendor", "portal"]):
            self._add_node(nodes, "partners", "Partner / External Channels", "channels")

        self._add_node(nodes, "frontend", "Frontend Apps / BFF", "experience")
        if self._mentions_any(text, ["aem", "content", "cms", "experience"]):
            self._add_node(nodes, "aem", "AEM", "experience")
        if self._mentions_any(text, ["aem assets", "assets", "dam"]):
            self._add_node(nodes, "aem_assets", "AEM Assets", "experience")
        if self._mentions_any(text, ["adobe commerce", "commerce", "storefront", "cart", "checkout"]):
            self._add_node(nodes, "commerce", "Adobe Commerce", "experience")
        if self._mentions_any(text, ["cdn", "edge", "front door"]):
            self._add_node(nodes, "edge", "Edge Delivery / CDN", "experience")
        if self._mentions_any(text, ["personalization", "personalisation", "targeting", "recommendation"]):
            self._add_node(nodes, "personalization", "Personalization Engine", "experience")
        if self._mentions_any(text, ["adobe target", "target"]):
            self._add_node(nodes, "target", "Adobe Target", "experience")
        if self._mentions_any(text, ["dynamic media", "adobe dynamic media"]):
            self._add_node(nodes, "dynamic_media", "Adobe Dynamic Media", "experience")
        if self._mentions_any(text, ["seo", "search engine", "discoverability", "metadata"]):
            self._add_node(nodes, "seo", "SEO / Metadata Services", "experience")
        if not any(node.key == "aem" for node in nodes):
            self._add_node(nodes, "portal", "Digital Experience Portal", "experience")

        self._add_node(nodes, "entra", self._cloud_identity_label(cloud_provider), "identity")
        if self._mentions_any(text, ["private", "network", "private endpoint", "private link", "vnet", "secure"]):
            self._add_node(nodes, "network", "Private Network Boundary", "identity")

        self._add_node(nodes, "apim", self._cloud_api_management_label(cloud_provider), "integration")
        self._add_node(nodes, "api_gateway", "API Gateway", "integration")
        if self._mentions_any(text, ["api gateway", "gateway", "edge api"]):
            self._add_node(nodes, "edge_api", "Edge API Policies", "integration")
        if self._mentions_any(text, ["event", "async", "messaging", "queue", "service bus"]):
            self._add_node(nodes, "service_bus", self._cloud_messaging_label(cloud_provider), "integration")
        if self._mentions_any(text, ["event grid", "pub/sub", "webhook"]):
            self._add_node(nodes, "event_grid", self._cloud_eventing_label(cloud_provider), "integration")
        self._add_node(nodes, "orchestration", "Integration Orchestration", "integration")

        if self._contains_any_service(analysis, ["Azure Container Apps", "AWS App Runner", "Cloud Run"]):
            self._add_node(nodes, "container_apps", self._cloud_container_runtime_label(cloud_provider), "applications")
        if self._contains_any_service(analysis, ["Azure Kubernetes Service", "Amazon EKS", "Google Kubernetes Engine"]):
            self._add_node(nodes, "aks", self._cloud_kubernetes_label(cloud_provider), "applications")
        self._add_node(nodes, "domain_services", "Domain Services", "applications")
        self._add_node(nodes, "microservices", "Microservices", "applications")
        if self._mentions_any(text, ["microservice", "micro-service", "bounded context"]):
            self._add_node(nodes, "service_mesh", "Service-to-Service Policies", "applications")
        if self._mentions_any(text, ["soa", "service-oriented", "service oriented", "shared service"]):
            self._add_node(nodes, "shared_services", "Shared Enterprise Services", "applications")
        if self._contains_any_service(analysis, ["Azure AI Foundry", "Amazon Bedrock", "Vertex AI"]):
            self._add_node(nodes, "ai_foundry", self._cloud_ai_platform_label(cloud_provider), "applications")

        if self._mentions_any(text, ["sap"]):
            self._add_node(nodes, "sap", self._extract_sap_label(text), "systems")
        if self._mentions_any(text, ["cpi", "sap cpi", "cloud platform integration"]):
            self._add_node(nodes, "sap_cpi", "SAP CPI", "systems")
        if self._mentions_any(text, ["btp", "sap btp", "business technology platform"]):
            self._add_node(nodes, "sap_btp", "SAP BTP", "systems")
        if self._mentions_any(text, ["fiori", "sap fiori"]):
            self._add_node(nodes, "fiori", "SAP Fiori", "systems")
        if self._mentions_any(text, ["salesforce"]):
            self._add_node(nodes, "salesforce", "Salesforce", "systems")
        if self._mentions_any(text, ["martech", "campaign", "journey orchestration", "marketing automation"]):
            self._add_node(nodes, "martech", "MarTech Platform", "systems")
        if self._mentions_any(text, ["loyalty", "rewards", "membership"]):
            self._add_node(nodes, "loyalty", "Loyalty Platform", "systems")
        if self._mentions_any(text, ["search", "site search", "product search", "seo"]):
            self._add_node(nodes, "search_platform", "Search Platform", "systems")
        if self._mentions_any(text, ["data platform", "lakehouse", "warehouse", "fabric", "databricks"]):
            self._add_node(nodes, "data_platform", "Enterprise Data Platform", "systems")
        if self._mentions_any(text, ["erp", "crm", "backend", "line of business"]) and len(
            [node for node in nodes if node.category == "systems"]
        ) < 2:
            self._add_node(nodes, "lob", "Line-of-Business Systems", "systems")

        if self._contains_any_service(
            analysis,
            ["Azure Storage or Cosmos DB", "Azure Storage", "Cosmos DB", "Amazon S3 or DynamoDB", "Amazon S3", "Amazon DynamoDB", "Cloud Storage or Firestore", "Cloud Storage", "Firestore"],
        ):
            self._add_node(nodes, "operational_data", "Operational Data Store", "data")
        self._add_node(nodes, "canonical_data", "Canonical Data Model", "data")
        if self._contains_any_service(
            analysis,
            ["Azure Data Lake / Synapse / Fabric", "Synapse", "Fabric", "Amazon S3, Redshift, Glue", "Amazon Redshift", "BigQuery, Dataplex, Cloud Storage", "BigQuery", "Dataplex"],
        ):
            self._add_node(nodes, "analytics", "Analytics & Reporting", "data")
        if self._mentions_any(text, ["personalization", "personalisation", "customer profile", "cdp"]):
            self._add_node(nodes, "customer_profile", "Customer Profile / CDP", "data")
        if self._mentions_any(text, ["adobe analytics"]):
            self._add_node(nodes, "adobe_analytics", "Adobe Analytics", "data")
        if self._mentions_any(text, ["aep", "adobe experience platform"]):
            self._add_node(nodes, "aep", "Adobe Experience Platform", "data")
        if self._mentions_any(text, ["rt-cdp", "rtcdp", "real-time cdp"]):
            self._add_node(nodes, "rt_cdp", "Adobe RT-CDP", "data")
        if cloud_provider == "azure" or self._mentions_any(text, ["azure"]):
            self._add_node(nodes, "azure_cloud", "Azure Platform", "data")
        if cloud_provider == "aws" or self._mentions_any(text, ["aws"]):
            self._add_node(nodes, "aws_cloud", "AWS Platform", "data")
        if cloud_provider == "gcp" or self._mentions_any(text, ["gcp", "google cloud"]):
            self._add_node(nodes, "gcp_cloud", "GCP Platform", "data")
        if self._contains_any_service(analysis, ["Azure AI Search", "Amazon OpenSearch Service", "Vertex AI Search"]) or self._mentions_any(text, ["seo", "search", "discoverability"]):
            self._add_node(nodes, "search_index", "Search / SEO Index", "data")

        self._add_node(nodes, "monitor", self._cloud_monitoring_label(cloud_provider), "observability")
        self._add_node(nodes, "governance", "Audit / Governance", "observability")
        self._add_runtime_nodes(nodes, analysis)

        return nodes

    def _build_edges(self, nodes: list[DiagramNode]) -> list[DiagramEdge]:
        node_map = {node.key: node for node in nodes}
        edges: list[DiagramEdge] = []

        def connect(source: str, target: str, label: str, style: str = "solid") -> None:
            if source in node_map and target in node_map:
                edge = DiagramEdge(source=source, target=target, label=label, style=style)
                if edge not in edges:
                    edges.append(edge)

        connect("users", "web", "Journeys")
        connect("partners", "web", "Partner Access")

        for experience in self._node_keys(nodes, "experience"):
            connect("web", experience, "UX / Content")
            connect("partners", experience, "Partner Experience")
            connect(experience, "entra", "SSO / Session")
            connect(experience, "apim", "Experience APIs")

        connect("web", "edge", "Global Delivery")
        connect("web", "frontend", "Channel Requests")
        connect("edge", "frontend", "Edge Routing")
        connect("aem_assets", "aem", "Asset Supply")
        connect("frontend", "apim", "BFF / App APIs")
        connect("frontend", "api_gateway", "Northbound APIs")
        connect("seo", "aem", "Structured Metadata")
        connect("seo", "search_index", "SEO Feeds")
        connect("personalization", "frontend", "Targeted Content")
        connect("target", "personalization", "Offer Decisions")
        connect("dynamic_media", "frontend", "Optimized Media")
        connect("personalization", "customer_profile", "Profile Decisions")

        connect("entra", "apim", "AuthN / Tokens")
        connect("network", "apim", "Private Access")
        connect("api_gateway", "apim", "North-South Policies")
        connect("edge_api", "api_gateway", "Edge Controls")
        connect("apim", "orchestration", "Policy / Routing")
        connect("event_grid", "orchestration", "Events / Notifications")
        connect("orchestration", "domain_services", "Service Composition")
        connect("orchestration", "microservices", "Domain APIs")
        connect("orchestration", "shared_services", "Shared Services")
        connect("orchestration", "service_bus", "Events")
        connect("service_bus", "domain_services", "Async Commands")
        connect("service_bus", "microservices", "Domain Events")
        connect("domain_services", "container_apps", "Runtime")
        connect("domain_services", "aks", "Runtime")
        connect("domain_services", "ai_foundry", "AI Enrichment")
        connect("domain_services", "operational_data", "Transactional Data")
        connect("domain_services", "canonical_data", "Canonical Events")
        connect("microservices", "container_apps", "Service Runtime")
        connect("microservices", "operational_data", "Service Data")
        connect("microservices", "canonical_data", "Business Events")
        connect("service_mesh", "microservices", "East-West Controls")
        connect("service_bus", "canonical_data", "Event Stream")

        for system in self._node_keys(nodes, "systems"):
            connect("orchestration", system, "API / Event Integration")
            connect(system, "canonical_data", "Master Data")

        connect("sap_btp", "sap", "Extensions")
        connect("sap_cpi", "sap", "Process Integration")
        connect("fiori", "sap", "Business UX")

        connect("operational_data", "analytics", "Batch / CDC")
        connect("canonical_data", "analytics", "Analytics Feed")
        connect("customer_profile", "analytics", "Customer Insights")
        connect("adobe_analytics", "analytics", "Digital Insights")
        connect("aep", "customer_profile", "Unified Profiles")
        connect("rt_cdp", "customer_profile", "Audience Activation")
        connect("canonical_data", "customer_profile", "Profile Inputs")
        connect("canonical_data", "search_index", "Searchable Content")
        connect("analytics", "web", "Insights / Reporting")
        connect("analytics", "personalization", "Segments / Models")
        connect("azure_cloud", "analytics", "Cloud Data Services")
        connect("aws_cloud", "search_index", "Cloud Support")
        connect("gcp_cloud", "analytics", "ML / Analytics")

        for observable in (
            "apim",
            "api_gateway",
            "orchestration",
            "domain_services",
            "microservices",
            "service_bus",
            "event_grid",
            "container_apps",
            "aks",
            "frontend",
        ):
            connect(observable, "monitor", "Telemetry", style="dashed")
        for governed in ("entra", "apim", "api_gateway", "canonical_data", "analytics", "customer_profile", "search_index"):
            connect(governed, "governance", "Audit", style="dashed")

        return edges

    def _build_drawio_xml(self, nodes: list[DiagramNode], edges: list[DiagramEdge]) -> str:
        node_map = {node.key: node for node in nodes}
        display_keys = [node.key for node in self._ordered_nodes(nodes)]
        geometry_map = self._layout_node_geometry(nodes)
        page_width = 2700
        page_height = 1450

        lines = [
            '<mxfile host="app.diagrams.net">',
            '  <diagram id="enterprise-architecture" name="Enterprise Architecture">',
            f'    <mxGraphModel dx="1800" dy="1080" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="{page_width}" pageHeight="{page_height}" math="0" shadow="0">',
            "      <root>",
            '        <mxCell id="0" />',
            '        <mxCell id="1" parent="0" />',
            '        <mxCell id="diagram-title" value="Enterprise Architecture Overview" style="text;html=1;strokeColor=none;fillColor=none;align=left;verticalAlign=middle;whiteSpace=wrap;fontSize=24;fontStyle=1;fontColor=#0f172a;" vertex="1" parent="1">',
            '          <mxGeometry x="40" y="24" width="520" height="34" as="geometry" />',
            '        </mxCell>',
            '        <mxCell id="diagram-subtitle" value="Clean solution view with primary user journeys and integration flows" style="text;html=1;strokeColor=none;fillColor=none;align=left;verticalAlign=middle;whiteSpace=wrap;fontSize=12;fontColor=#64748b;" vertex="1" parent="1">',
            '          <mxGeometry x="40" y="54" width="640" height="22" as="geometry" />',
            '        </mxCell>',
        ]

        for heading, x, y in self._drawio_section_headings():
            heading_id = self._drawio_safe_id(heading)
            lines.extend(
                [
                    f'        <mxCell id="heading-{heading_id}" value="{escape(heading)}" style="text;html=1;strokeColor=none;fillColor=none;align=left;verticalAlign=middle;whiteSpace=wrap;fontSize=14;fontStyle=1;fontColor=#334155;" vertex="1" parent="1">',
                    f'          <mxGeometry x="{x}" y="{y}" width="220" height="24" as="geometry" />',
                    "        </mxCell>",
                ]
            )

        node_ids: dict[str, str] = {}
        for key in display_keys:
            node = node_map[key]
            x, y, width, height = geometry_map[key]

            if key in {"users", "partners"}:
                cell_id = f"node-{key}"
                node_ids[key] = cell_id
                lines.extend(
                    [
                        f'        <mxCell id="{cell_id}" value="" style="shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top;html=1;fillColor=#ffffff;strokeColor=#475569;strokeWidth=1.5;" vertex="1" parent="1">',
                        f'          <mxGeometry x="{x}" y="{y}" width="{width}" height="{height}" as="geometry" />',
                        "        </mxCell>",
                        f'        <mxCell id="{cell_id}-label" value="{escape(node.label)}" style="text;html=1;strokeColor=none;fillColor=none;align=center;verticalAlign=top;whiteSpace=wrap;fontSize=12;fontStyle=1;fontColor=#0f172a;" vertex="1" parent="1">',
                        f'          <mxGeometry x="{x - 28}" y="{y + height + 8}" width="{width + 56}" height="20" as="geometry" />',
                        "        </mxCell>",
                    ]
                )
                continue

            group_id = f"group-{key}"
            body_id = f"node-{key}"
            node_ids[key] = body_id
            lines.extend(self._build_drawio_grouped_node(group_id, body_id, node, x, y, width, height))

        edge_index = 1
        for edge in edges:
            source = node_ids.get(edge.source)
            target = node_ids.get(edge.target)
            if not source or not target:
                continue
            edge_style = (
                "edgeStyle=orthogonalEdgeStyle;rounded=1;orthogonalLoop=1;jettySize=auto;"
                "html=1;endArrow=block;strokeWidth=1.6;"
            )
            edge_style += "dashed=1;strokeColor=#94a3b8;" if edge.style == "dashed" else "strokeColor=#475569;"
            value = escape(edge.label) if edge.label else ""
            if value:
                edge_style += "fontSize=11;fontColor=#334155;labelBackgroundColor=#ffffff;"
            lines.extend(
                [
                    f'        <mxCell id="edge-{edge_index}" value="{value}" style="{edge_style}" edge="1" parent="1" source="{source}" target="{target}">',
                    '          <mxGeometry relative="1" as="geometry" />',
                    "        </mxCell>",
                ]
            )
            edge_index += 1

        lines.extend(
            [
                "      </root>",
                "    </mxGraphModel>",
                "  </diagram>",
                "</mxfile>",
            ]
        )
        return "\n".join(lines)

    def _add_node(self, nodes: list[DiagramNode], key: str, label: str, category: str) -> None:
        if any(node.key == key for node in nodes):
            return
        nodes.append(DiagramNode(key=key, label=label, category=category))

    def _add_runtime_nodes(self, nodes: list[DiagramNode], analysis: ArchitectureAnalysis) -> None:
        existing_labels = {node.label.casefold() for node in nodes}
        for raw_label in [*analysis.tools_and_technologies, *analysis.suggested_azure_services]:
            label = re.sub(r"\s+", " ", (raw_label or "").strip())
            if not label or label.casefold() in existing_labels:
                continue
            category = self._categorize_runtime_label(label)
            key = self._runtime_node_key(label, category, nodes)
            self._add_node(nodes, key, label, category)
            existing_labels.add(label.casefold())

    def _categorize_runtime_label(self, label: str) -> str:
        lowered = label.casefold()
        if any(token in lowered for token in ("identity", "iam", "entra", "auth", "sso", "key vault", "secret manager")):
            return "identity"
        if any(token in lowered for token in ("api", "gateway", "event", "pub/sub", "queue", "sns", "sqs", "service bus", "integration", "apigee")):
            return "integration"
        if any(token in lowered for token in ("kubernetes", "aks", "eks", "gke", "container", "run", "app runner", "bedrock", "vertex ai", "ai foundry")):
            return "applications"
        if any(token in lowered for token in ("storage", "database", "cosmos", "dynamodb", "firestore", "bigquery", "redshift", "fabric", "synapse", "analytics", "search", "opensearch")):
            return "data"
        if any(token in lowered for token in ("aem", "adobe", "commerce", "target", "cdn", "front door", "cloudfront")):
            return "experience"
        if any(token in lowered for token in ("sap", "salesforce", "crm", "erp", "platform")):
            return "systems"
        if any(token in lowered for token in ("monitor", "logging", "log analytics", "cloudwatch", "observability")):
            return "observability"
        return "applications"

    def _runtime_node_key(self, label: str, category: str, nodes: list[DiagramNode]) -> str:
        base = re.sub(r"[^a-z0-9]+", "_", label.casefold()).strip("_") or category
        key = f"{category}_{base}"
        suffix = 2
        existing = {node.key for node in nodes}
        while key in existing:
            key = f"{category}_{base}_{suffix}"
            suffix += 1
        return key

    def _node_keys(self, nodes: list[DiagramNode], category: str) -> list[str]:
        return [node.key for node in nodes if node.category == category]

    def _ordered_nodes(self, nodes: list[DiagramNode]) -> list[DiagramNode]:
        category_rank = {category: index for index, category in enumerate(self.CATEGORY_ORDER)}
        return sorted(nodes, key=lambda node: (category_rank.get(node.category, 999), node.label.casefold()))

    def _layout_node_geometry(self, nodes: list[DiagramNode]) -> dict[str, tuple[int, int, int, int]]:
        category_positions = {
            "channels": (70, 180),
            "experience": (360, 180),
            "identity": (710, 180),
            "integration": (970, 180),
            "applications": (1280, 180),
            "systems": (1640, 180),
            "data": (1080, 760),
            "observability": (1710, 760),
        }
        category_widths = {
            "channels": 220,
            "experience": 260,
            "identity": 220,
            "integration": 250,
            "applications": 300,
            "systems": 280,
            "data": 280,
            "observability": 230,
        }
        category_heights = {
            "channels": 90,
            "experience": 126,
            "identity": 110,
            "integration": 120,
            "applications": 120,
            "systems": 120,
            "data": 110,
            "observability": 90,
        }
        geometry: dict[str, tuple[int, int, int, int]] = {}
        category_counts: dict[str, int] = {category: 0 for category in self.CATEGORY_ORDER}
        for node in self._ordered_nodes(nodes):
            count = category_counts.get(node.category, 0)
            base_x, base_y = category_positions.get(node.category, (150, 200))
            width = category_widths.get(node.category, 220)
            height = category_heights.get(node.category, 110)
            columns = 2 if node.category in {"experience", "applications", "systems", "data"} else 1
            column = count % columns
            row = count // columns
            x = base_x + (column * (width + 24))
            y = base_y + (row * (height + 32))
            if node.key in {"users", "partners"}:
                width = 46
                height = 82
            geometry[node.key] = (x, y, width, height)
            category_counts[node.category] = count + 1
        return geometry

    def _contains_service(self, analysis: ArchitectureAnalysis, service_name: str) -> bool:
        return any(service_name in service for service in analysis.suggested_azure_services)

    def _contains_any_service(self, analysis: ArchitectureAnalysis, service_names: list[str]) -> bool:
        return any(name in service for service in analysis.suggested_azure_services for name in service_names)

    def _cloud_identity_label(self, cloud_provider: str) -> str:
        return {
            "aws": "AWS IAM Identity Center",
            "gcp": "Cloud Identity",
        }.get(cloud_provider, "Microsoft Entra ID")

    def _cloud_api_management_label(self, cloud_provider: str) -> str:
        return {
            "aws": "Amazon API Gateway",
            "gcp": "Apigee",
        }.get(cloud_provider, "Azure API Management")

    def _cloud_messaging_label(self, cloud_provider: str) -> str:
        return {
            "aws": "Amazon SQS / SNS",
            "gcp": "Pub/Sub",
        }.get(cloud_provider, "Azure Service Bus")

    def _cloud_eventing_label(self, cloud_provider: str) -> str:
        return {
            "aws": "Amazon EventBridge",
            "gcp": "Eventarc",
        }.get(cloud_provider, "Azure Event Grid")

    def _cloud_container_runtime_label(self, cloud_provider: str) -> str:
        return {
            "aws": "AWS App Runner",
            "gcp": "Cloud Run",
        }.get(cloud_provider, "Azure Container Apps")

    def _cloud_kubernetes_label(self, cloud_provider: str) -> str:
        return {
            "aws": "Amazon EKS",
            "gcp": "Google Kubernetes Engine",
        }.get(cloud_provider, "Azure Kubernetes Service")

    def _cloud_ai_platform_label(self, cloud_provider: str) -> str:
        return {
            "aws": "Amazon Bedrock",
            "gcp": "Vertex AI",
        }.get(cloud_provider, "Azure AI Foundry")

    def _cloud_monitoring_label(self, cloud_provider: str) -> str:
        return {
            "aws": "Amazon CloudWatch",
            "gcp": "Cloud Monitoring",
        }.get(cloud_provider, "Azure Monitor")

    def _mentions_any(self, text: str, terms: list[str]) -> bool:
        lowered = text.casefold()
        return any(term.casefold() in lowered for term in terms)

    def _extract_sap_label(self, text: str) -> str:
        match = re.search(r"\bSAP(?:\s+[A-Za-z0-9/.-]+)?\b", text, flags=re.IGNORECASE)
        if not match:
            return "SAP"
        label = re.sub(r"\s+", " ", match.group(0)).strip()
        return label.upper() if label.casefold() == "sap" else label.replace("Sap", "SAP")

    def _build_drawio_node_value(self, node: DiagramNode) -> str:
        details = self._node_details(node.key, node.label)
        if not details:
            return node.label

        detail_html = "".join(
            f"<div style='font-size:11px;color:#475569;line-height:1.35;'>{escape(detail)}</div>" for detail in details[:3]
        )
        return (
            f"<div style='font-size:15px;font-weight:700;color:#0f172a;line-height:1.25;margin-bottom:6px;'>"
            f"{escape(node.label)}"
            "</div>"
            f"{detail_html}"
        )

    def _should_label_drawio_edge(self, edge: DiagramEdge) -> bool:
        return (edge.source, edge.target) in self.DRAWIO_PRIMARY_EDGE_LABELS

    def _build_drawio_compact_value(self, node: DiagramNode) -> str:
        details = self._node_details(node.key, node.label)
        subtitle = " | ".join(details[:2])
        if not subtitle:
            return f"<div style='font-size:15px;font-weight:700;color:#0f172a;'>{escape(node.label)}</div>"
        return (
            f"<div style='font-size:15px;font-weight:700;color:#0f172a;line-height:1.2;margin-bottom:6px;'>{escape(node.label)}</div>"
            f"<div style='font-size:11px;color:#64748b;line-height:1.3;'>{escape(subtitle)}</div>"
        )

    def _drawio_display_node_keys(self) -> list[str]:
        return [
            "users",
            "partners",
            "web",
            "edge",
            "frontend",
            "aem",
            "aem_assets",
            "commerce",
            "personalization",
            "target",
            "dynamic_media",
            "entra",
            "network",
            "apim",
            "orchestration",
            "service_bus",
            "domain_services",
            "microservices",
            "ai_foundry",
            "sap",
            "sap_cpi",
            "sap_btp",
            "fiori",
            "salesforce",
            "martech",
            "canonical_data",
            "customer_profile",
            "search_index",
            "analytics",
            "adobe_analytics",
            "aep",
            "rt_cdp",
            "azure_cloud",
            "aws_cloud",
            "gcp_cloud",
            "monitor",
            "governance",
        ]

    def _drawio_section_headings(self) -> list[tuple[str, int, int]]:
        return [
            ("Users", 40, 108),
            ("Channels", 210, 108),
            ("Experience", 430, 108),
            ("Security", 980, 108),
            ("Integration", 1230, 108),
            ("Applications", 1760, 108),
            ("Enterprise Systems", 2060, 108),
            ("Data & Insights", 1280, 700),
            ("Operations", 1870, 610),
        ]

    def _drawio_node_geometry(self, key: str) -> tuple[int, int, int, int]:
        positions = {
            "users": (60, 170, 42, 80),
            "partners": (60, 340, 42, 80),
            "web": (190, 210, 190, 96),
            "edge": (430, 110, 210, 132),
            "frontend": (430, 300, 210, 132),
            "aem": (700, 80, 250, 188),
            "aem_assets": (700, 300, 250, 150),
            "commerce": (700, 500, 250, 168),
            "personalization": (700, 720, 250, 168),
            "target": (990, 500, 180, 132),
            "dynamic_media": (990, 300, 180, 132),
            "entra": (980, 90, 190, 114),
            "network": (980, 230, 190, 96),
            "apim": (1230, 90, 220, 132),
            "orchestration": (1490, 90, 230, 132),
            "service_bus": (1490, 280, 230, 114),
            "domain_services": (1760, 90, 240, 132),
            "microservices": (1760, 280, 240, 132),
            "ai_foundry": (1760, 470, 240, 114),
            "sap": (2060, 70, 250, 188),
            "sap_cpi": (2060, 290, 220, 114),
            "sap_btp": (2060, 430, 220, 114),
            "fiori": (2060, 570, 220, 114),
            "salesforce": (2330, 290, 200, 114),
            "martech": (2330, 430, 200, 114),
            "canonical_data": (1280, 760, 240, 114),
            "customer_profile": (1560, 760, 240, 114),
            "search_index": (1840, 760, 240, 114),
            "analytics": (2120, 760, 240, 114),
            "adobe_analytics": (1280, 920, 220, 114),
            "aep": (1540, 920, 220, 114),
            "rt_cdp": (1800, 920, 220, 114),
            "azure_cloud": (2060, 920, 220, 114),
            "aws_cloud": (2320, 1080, 200, 114),
            "gcp_cloud": (2320, 920, 200, 114),
            "monitor": (2120, 1110, 180, 78),
            "governance": (2120, 1210, 180, 78),
        }
        return positions.get(key, (210, 230, 200, 76))

    def _drawio_edge_specs(self, node_ids: dict[str, str]) -> list[tuple[str, str, str, str]]:
        preferred_edges = [
            ("users", "web", "User journeys", "solid"),
            ("partners", "web", "Partner access", "solid"),
            ("web", "frontend", "Channel requests", "solid"),
            ("edge", "frontend", "Accelerated delivery", "solid"),
            ("aem", "frontend", "Content delivery", "solid"),
            ("aem_assets", "aem", "Asset supply", "solid"),
            ("commerce", "frontend", "Commerce services", "solid"),
            ("personalization", "frontend", "Personalized experiences", "solid"),
            ("target", "personalization", "Targeting decisions", "solid"),
            ("dynamic_media", "frontend", "Dynamic media", "solid"),
            ("frontend", "entra", "Login / session", "solid"),
            ("entra", "apim", "Identity tokens", "solid"),
            ("frontend", "apim", "Experience APIs", "solid"),
            ("apim", "orchestration", "Managed integration", "solid"),
            ("orchestration", "service_bus", "Events", "solid"),
            ("orchestration", "domain_services", "Business orchestration", "solid"),
            ("service_bus", "microservices", "Async events", "solid"),
            ("domain_services", "sap", "ERP integration", "solid"),
            ("sap_cpi", "sap", "SAP integration", "solid"),
            ("sap_btp", "sap", "Extensions", "solid"),
            ("fiori", "sap", "User experience", "solid"),
            ("domain_services", "salesforce", "CRM integration", "solid"),
            ("domain_services", "martech", "Campaign activation", "solid"),
            ("domain_services", "canonical_data", "Business events", "solid"),
            ("canonical_data", "customer_profile", "Profile inputs", "solid"),
            ("canonical_data", "search_index", "Content indexing", "solid"),
            ("canonical_data", "analytics", "Analytics feed", "solid"),
            ("adobe_analytics", "analytics", "Adobe metrics", "solid"),
            ("aep", "customer_profile", "Experience profiles", "solid"),
            ("rt_cdp", "customer_profile", "Audience activation", "solid"),
            ("analytics", "personalization", "Segments / models", "solid"),
            ("apim", "monitor", "Telemetry", "dashed"),
            ("canonical_data", "governance", "Audit", "dashed"),
        ]

        if "azure_cloud" in node_ids:
            preferred_edges.append(("azure_cloud", "analytics", "Azure data platform", "solid"))
        if "aws_cloud" in node_ids:
            preferred_edges.append(("aws_cloud", "analytics", "AWS data platform", "solid"))
        if "gcp_cloud" in node_ids:
            preferred_edges.append(("gcp_cloud", "analytics", "GCP analytics", "solid"))

        result: list[tuple[str, str, str, str]] = []
        for source_key, target_key, label, style in preferred_edges:
            if source_key in node_ids and target_key in node_ids:
                result.append((node_ids[source_key], node_ids[target_key], label, style))
        return result

    def _build_drawio_grouped_node(
        self,
        group_id: str,
        body_id: str,
        node: DiagramNode,
        x: int,
        y: int,
        width: int,
        height: int,
    ) -> list[str]:
        lines = [
            f'        <mxCell id="{group_id}" value="" style="group" vertex="1" connectable="0" parent="1">',
            f'          <mxGeometry x="{x}" y="{y}" width="{width}" height="{height}" as="geometry" />',
            "        </mxCell>",
        ]

        header_height = 42
        details = self._node_details(node.key, node.label)
        body_style = self._drawio_group_body_style(node.category)
        tile_style = self._drawio_group_tile_style(node.category)

        lines.extend(
            [
                f'        <mxCell id="{body_id}" value="" style="{body_style}" vertex="1" parent="{group_id}">',
                f'          <mxGeometry x="0" y="0" width="{width}" height="{height}" as="geometry" />',
                "        </mxCell>",
                f'        <mxCell id="{body_id}-title" value="{escape(node.label)}" style="text;html=1;strokeColor=none;fillColor=none;align=center;verticalAlign=middle;whiteSpace=wrap;fontSize=16;fontStyle=1;fontColor=#0f172a;" vertex="1" parent="{group_id}">',
                f'          <mxGeometry x="16" y="10" width="{width - 32}" height="24" as="geometry" />',
                "        </mxCell>",
            ]
        )

        if not details:
            return lines

        tile_gap = 10
        tile_y_start = header_height
        tile_width = (width - 30) / 2
        tile_height = 28
        max_rows = max(1, int((height - tile_y_start - 10 + 8) // (tile_height + 8)))
        max_tiles = max_rows * 2
        details = details[:max_tiles]

        for index, detail in enumerate(details):
            col = index % 2
            row = index // 2
            tile_x = 10 + (col * (tile_width + tile_gap))
            tile_y = tile_y_start + (row * (tile_height + 8))
            lines.extend(
                [
                    f'        <mxCell id="{body_id}-detail-{index}" value="{escape(detail)}" style="{tile_style}" vertex="1" parent="{group_id}">',
                    f'          <mxGeometry x="{tile_x}" y="{tile_y}" width="{tile_width}" height="{tile_height}" as="geometry" />',
                    "        </mxCell>",
                ]
            )

        return lines

    def _drawio_group_body_style(self, category: str) -> str:
        category_styles = {
            "channels": "rounded=1;whiteSpace=wrap;html=1;fillColor=#f2f2f2;strokeColor=#cbd5e1;shadow=0;",
            "experience": "rounded=1;whiteSpace=wrap;html=1;fillColor=#fff2cc;strokeColor=#d6b656;shadow=0;",
            "identity": "rounded=1;whiteSpace=wrap;html=1;fillColor=#fff2cc;strokeColor=#d6b656;shadow=0;",
            "integration": "rounded=1;whiteSpace=wrap;html=1;fillColor=#dae8fc;strokeColor=#6c8ebf;shadow=0;",
            "applications": "rounded=1;whiteSpace=wrap;html=1;fillColor=#d5e8d4;strokeColor=#82b366;shadow=0;",
            "systems": "rounded=1;whiteSpace=wrap;html=1;fillColor=#f8cecc;strokeColor=#b85450;shadow=0;",
            "data": "rounded=1;whiteSpace=wrap;html=1;fillColor=#e1d5e7;strokeColor=#9673a6;shadow=0;",
            "observability": "rounded=1;whiteSpace=wrap;html=1;fillColor=#f5f5f5;strokeColor=#666666;shadow=0;",
        }
        return category_styles[category]

    def _drawio_group_tile_style(self, category: str) -> str:
        tile_styles = {
            "channels": "rounded=1;whiteSpace=wrap;html=1;fillColor=#d8d8d8;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
            "experience": "rounded=1;whiteSpace=wrap;html=1;fillColor=#fee599;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
            "identity": "rounded=1;whiteSpace=wrap;html=1;fillColor=#fee599;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
            "integration": "rounded=1;whiteSpace=wrap;html=1;fillColor=#9fc5f8;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
            "applications": "rounded=1;whiteSpace=wrap;html=1;fillColor=#b6d7a8;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
            "systems": "rounded=1;whiteSpace=wrap;html=1;fillColor=#ea9999;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
            "data": "rounded=1;whiteSpace=wrap;html=1;fillColor=#d5a6bd;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
            "observability": "rounded=1;whiteSpace=wrap;html=1;fillColor=#d9d9d9;strokeColor=none;align=center;verticalAlign=middle;fontSize=12;fontColor=#111827;",
        }
        return tile_styles[category]

    def _drawio_safe_id(self, value: str) -> str:
        normalized = re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")
        return normalized or "section"

    def _drawio_edge_anchor_style(self, source: str, target: str, node_ids: dict[str, str]) -> str:
        node_keys = {value: key for key, value in node_ids.items()}
        source_key = node_keys.get(source)
        target_key = node_keys.get(target)
        anchor_map = {
            ("users", "web"): "exitX=1;exitY=0.5;entryX=0;entryY=0.35;",
            ("partners", "web"): "exitX=1;exitY=0.5;entryX=0;entryY=0.75;",
            ("web", "frontend"): "exitX=1;exitY=0.7;entryX=0;entryY=0.35;",
            ("edge", "frontend"): "exitX=0.5;exitY=1;entryX=0.35;entryY=0;",
            ("aem", "frontend"): "exitX=0;exitY=0.6;entryX=1;entryY=0.2;",
            ("aem_assets", "aem"): "exitX=0.5;exitY=0;entryX=0.5;entryY=1;",
            ("commerce", "frontend"): "exitX=0;exitY=0.2;entryX=0.5;entryY=1;",
            ("personalization", "frontend"): "exitX=0;exitY=0.5;entryX=0.5;entryY=1;",
            ("target", "personalization"): "exitX=0;exitY=0.5;entryX=1;entryY=0.35;",
            ("dynamic_media", "frontend"): "exitX=0;exitY=0.5;entryX=1;entryY=0.65;",
            ("frontend", "entra"): "exitX=1;exitY=0.1;entryX=0;entryY=0.5;",
            ("entra", "apim"): "exitX=1;exitY=0.5;entryX=0;entryY=0.25;",
            ("frontend", "apim"): "exitX=1;exitY=0.5;entryX=0;entryY=0.7;",
            ("apim", "orchestration"): "exitX=1;exitY=0.5;entryX=0;entryY=0.5;",
            ("orchestration", "service_bus"): "exitX=0.5;exitY=1;entryX=0.5;entryY=0;",
            ("orchestration", "domain_services"): "exitX=1;exitY=0.45;entryX=0;entryY=0.45;",
            ("service_bus", "microservices"): "exitX=1;exitY=0.5;entryX=0;entryY=0.5;",
            ("domain_services", "sap"): "exitX=1;exitY=0.25;entryX=0;entryY=0.45;",
            ("domain_services", "salesforce"): "exitX=1;exitY=0.5;entryX=0;entryY=0.5;",
            ("domain_services", "martech"): "exitX=1;exitY=0.8;entryX=0;entryY=0.5;",
            ("domain_services", "canonical_data"): "exitX=0.4;exitY=1;entryX=0.5;entryY=0;",
            ("canonical_data", "customer_profile"): "exitX=1;exitY=0.5;entryX=0;entryY=0.5;",
            ("canonical_data", "search_index"): "exitX=1;exitY=0.75;entryX=0;entryY=0.25;",
            ("canonical_data", "analytics"): "exitX=1;exitY=0.3;entryX=0;entryY=0.3;",
            ("adobe_analytics", "analytics"): "exitX=1;exitY=0.3;entryX=0.2;entryY=1;",
            ("aep", "customer_profile"): "exitX=0.5;exitY=0;entryX=0.5;entryY=1;",
            ("rt_cdp", "customer_profile"): "exitX=0.5;exitY=0;entryX=0.8;entryY=1;",
            ("azure_cloud", "analytics"): "exitX=0.5;exitY=0;entryX=0.8;entryY=1;",
            ("aws_cloud", "analytics"): "exitX=0;exitY=0.4;entryX=1;entryY=0.8;",
            ("gcp_cloud", "analytics"): "exitX=0;exitY=0.4;entryX=1;entryY=0.8;",
            ("analytics", "personalization"): "exitX=0;exitY=0.15;entryX=1;entryY=0.85;",
            ("apim", "monitor"): "exitX=0.7;exitY=1;entryX=0.5;entryY=0;",
            ("canonical_data", "governance"): "exitX=1;exitY=1;entryX=0;entryY=0.3;",
        }
        return anchor_map.get((source_key, target_key), "")

    def _node_details(self, key: str, label: str) -> list[str]:
        cloud_specific = {
            "entra": {
                "AWS IAM Identity Center": ["SSO", "SAML / OIDC", "Role-based access"],
                "Cloud Identity": ["SSO", "OIDC / SAML", "Role-based access"],
                "Microsoft Entra ID": ["SSO", "OAuth / OIDC", "Role-based access"],
            },
            "apim": {
                "Amazon API Gateway": ["Policies", "Developer access", "Security enforcement"],
                "Apigee": ["Policies", "Developer portal", "Security enforcement"],
                "Azure API Management": ["Policies", "Developer access", "Security enforcement"],
            },
            "service_bus": {
                "Amazon SQS / SNS": ["Queues", "Topics", "Reliable delivery"],
                "Pub/Sub": ["Topics", "Subscriptions", "Reliable delivery"],
                "Azure Service Bus": ["Queues", "Topics", "Reliable delivery"],
            },
            "event_grid": {
                "Amazon EventBridge": ["Event routing", "Webhook fan-out", "Notifications"],
                "Eventarc": ["Event routing", "Triggers", "Notifications"],
                "Azure Event Grid": ["Event distribution", "Webhook fan-out", "Notifications"],
            },
            "container_apps": {
                "AWS App Runner": ["Scalable runtime", "Managed revisions", "Service endpoints"],
                "Cloud Run": ["Serverless runtime", "Managed revisions", "Service endpoints"],
                "Azure Container Apps": ["Scalable runtime", "Revision rollout", "Service endpoints"],
            },
            "aks": {
                "Amazon EKS": ["Kubernetes runtime", "Platform controls", "Advanced workloads"],
                "Google Kubernetes Engine": ["Kubernetes runtime", "Platform controls", "Advanced workloads"],
                "Azure Kubernetes Service": ["Kubernetes runtime", "Platform controls", "Advanced workloads"],
            },
            "ai_foundry": {
                "Amazon Bedrock": ["Foundation models", "Guardrails", "AI orchestration"],
                "Vertex AI": ["Models", "Pipelines", "AI orchestration"],
                "Azure AI Foundry": ["Prompt flows", "Model access", "AI orchestration"],
            },
            "monitor": {
                "Amazon CloudWatch": ["Metrics", "Logs", "Alerting"],
                "Cloud Monitoring": ["Metrics", "Tracing", "Alerting"],
                "Azure Monitor": ["Metrics", "Tracing", "Alerting"],
            },
            "azure_cloud": {"Azure Platform": ["App services", "Integration", "Data platform", "Identity"]},
            "aws_cloud": {"AWS Platform": ["S3", "CloudFront", "Compute", "Messaging"]},
            "gcp_cloud": {"GCP Platform": ["BigQuery", "Vertex AI", "Storage", "Data processing"]},
        }
        if key in cloud_specific:
            return cloud_specific[key].get(label, self.DRAWIO_NODE_DETAILS.get(key, []))
        if key in self.DRAWIO_NODE_DETAILS:
            return self.DRAWIO_NODE_DETAILS.get(key, [])
        lowered = label.casefold()
        if any(token in lowered for token in ("api", "gateway", "apigee", "integration", "event", "queue", "topic")):
            return ["Contracts", "Security", "Routing"]
        if any(token in lowered for token in ("storage", "database", "analytics", "search", "index", "lake", "warehouse")):
            return ["Data store", "Query access", "Governance"]
        if any(token in lowered for token in ("kubernetes", "container", "run", "app runner", "compute", "bedrock", "vertex")):
            return ["Runtime", "Scaling", "Operations"]
        if any(token in lowered for token in ("identity", "auth", "entra", "iam", "secret", "vault")):
            return ["Identity", "Access control", "Secrets"]
        if any(token in lowered for token in ("commerce", "adobe", "cms", "experience", "target")):
            return ["Business capability", "Customer journeys", "Integration points"]
        return ["Core capability", "Enterprise integration", "Operations"]

    def _build_pptx_base64(self, nodes: list[DiagramNode], edges: list[DiagramEdge]) -> str:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        title = slide.shapes.add_textbox(Inches(0.45), Inches(0.2), Inches(6.2), Inches(0.45))
        title_tf = title.text_frame
        title_tf.text = "Enterprise Architecture Overview"
        title_run = title_tf.paragraphs[0].runs[0]
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(15, 23, 42)

        subtitle = slide.shapes.add_textbox(Inches(0.45), Inches(0.58), Inches(7.6), Inches(0.3))
        subtitle_tf = subtitle.text_frame
        subtitle_tf.text = "Native PowerPoint architecture blueprint aligned to the draw.io solution layout"
        subtitle_run = subtitle_tf.paragraphs[0].runs[0]
        subtitle_run.font.size = Pt(10)
        subtitle_run.font.color.rgb = RGBColor(100, 116, 139)

        accent_bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(0.45),
            Inches(0.98),
            Inches(12.25),
            Inches(0.05),
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = RGBColor(51, 92, 255)
        accent_bar.line.fill.background()
        canvas = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(0.42),
            Inches(1.2),
            Inches(12.45),
            Inches(5.9),
        )
        canvas.fill.solid()
        canvas.fill.fore_color.rgb = RGBColor(255, 255, 255)
        canvas.line.color.rgb = RGBColor(226, 232, 240)
        canvas.line.width = Pt(1)

        page_width = 2700
        page_height = 1450
        slide_width = 13.333
        slide_height = 7.5
        scale_x = slide_width / page_width
        scale_y = (slide_height - 1.12) / page_height
        x_offset = 0.52
        y_offset = 1.26
        node_map = {node.key: node for node in nodes}
        display_keys = [node.key for node in self._ordered_nodes(nodes)]
        geometry_map = self._layout_node_geometry(nodes)
        ppt_shapes: dict[str, object] = {}

        for heading, x, y in self._drawio_section_headings():
            heading_box = slide.shapes.add_textbox(
                Inches(x_offset + (x * scale_x)),
                Inches(y_offset + (y * scale_y) - 0.16),
                Inches(1.4),
                Inches(0.18),
            )
            heading_tf = heading_box.text_frame
            heading_tf.text = heading
            run = heading_tf.paragraphs[0].runs[0]
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)

        for key in display_keys:
            node = node_map[key]
            x, y, width, height = geometry_map[key]
            left = x_offset + (x * scale_x)
            top = y_offset + (y * scale_y)
            scaled_width = max(0.38, width * scale_x)
            scaled_height = max(0.32, height * scale_y)

            if key in {"users", "partners"}:
                head = slide.shapes.add_shape(
                    MSO_AUTO_SHAPE_TYPE.OVAL,
                    Inches(left + (scaled_width * 0.2)),
                    Inches(top),
                    Inches(max(0.12, scaled_width * 0.6)),
                    Inches(max(0.12, scaled_width * 0.6)),
                )
                head.fill.solid()
                head.fill.fore_color.rgb = RGBColor(255, 255, 255)
                head.line.color.rgb = RGBColor(71, 85, 105)
                head.line.width = Pt(1)

                body = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    int(Inches(left + (scaled_width / 2))),
                    int(Inches(top + (scaled_width * 0.6))),
                    int(Inches(left + (scaled_width / 2))),
                    int(Inches(top + scaled_height * 0.7)),
                )
                body.line.color.rgb = RGBColor(71, 85, 105)
                body.line.width = Pt(1)

                arms = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    int(Inches(left + scaled_width * 0.12)),
                    int(Inches(top + scaled_height * 0.42)),
                    int(Inches(left + scaled_width * 0.88)),
                    int(Inches(top + scaled_height * 0.42)),
                )
                arms.line.color.rgb = RGBColor(71, 85, 105)
                arms.line.width = Pt(1)

                leg_left = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    int(Inches(left + (scaled_width / 2))),
                    int(Inches(top + scaled_height * 0.7)),
                    int(Inches(left + scaled_width * 0.18)),
                    int(Inches(top + scaled_height)),
                )
                leg_left.line.color.rgb = RGBColor(71, 85, 105)
                leg_left.line.width = Pt(1)

                leg_right = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    int(Inches(left + (scaled_width / 2))),
                    int(Inches(top + scaled_height * 0.7)),
                    int(Inches(left + scaled_width * 0.82)),
                    int(Inches(top + scaled_height)),
                )
                leg_right.line.color.rgb = RGBColor(71, 85, 105)
                leg_right.line.width = Pt(1)

                title_box = slide.shapes.add_textbox(
                    Inches(max(0.1, left - 0.25)),
                    Inches(top + scaled_height + 0.06),
                    Inches(scaled_width + 0.5),
                    Inches(0.24),
                )
                title_tf = title_box.text_frame
                title_tf.clear()
                title_p = title_tf.paragraphs[0]
                title_p.alignment = PP_ALIGN.CENTER
                title_run = title_p.add_run()
                title_run.text = node.label
                title_run.font.size = Pt(8.5)
                title_run.font.bold = True
                title_run.font.color.rgb = RGBColor(15, 23, 42)

                ppt_shapes[key] = {
                    "left": left,
                    "top": top,
                    "width": scaled_width,
                    "height": scaled_height,
                }
                continue

            outer = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                Inches(left),
                Inches(top),
                Inches(scaled_width),
                Inches(scaled_height),
            )
            outer.fill.solid()
            outer.fill.fore_color.rgb = RGBColor.from_string(self.PPTX_NODE_FILL[node.category])
            outer.line.color.rgb = RGBColor.from_string(self.PPTX_NODE_LINE[node.category])
            outer.line.width = Pt(1.1)

            title_box = slide.shapes.add_textbox(
                Inches(left + 0.05),
                Inches(top + 0.04),
                Inches(max(0.25, scaled_width - 0.1)),
                Inches(0.2),
            )
            title_tf = title_box.text_frame
            title_tf.clear()
            title_p = title_tf.paragraphs[0]
            title_p.alignment = PP_ALIGN.CENTER
            title_run = title_p.add_run()
            title_run.text = node.label
            title_run.font.size = Pt(8.8)
            title_run.font.bold = True
            title_run.font.color.rgb = RGBColor(15, 23, 42)

            details = self._node_details(key, node.label)[:4]
            if details:
                detail_columns = 2 if scaled_width > 0.95 else 1
                tile_gap = 0.05
                inner_left = left + 0.05
                inner_top = top + 0.26
                tile_width = (scaled_width - 0.1 - ((detail_columns - 1) * tile_gap)) / detail_columns
                tile_height = 0.16
                for index, detail in enumerate(details):
                    col = index % detail_columns
                    row = index // detail_columns
                    tile_left = inner_left + (col * (tile_width + tile_gap))
                    tile_top = inner_top + (row * (tile_height + 0.04))
                    if tile_top + tile_height > top + scaled_height - 0.03:
                        break
                    tile = slide.shapes.add_shape(
                        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                        Inches(tile_left),
                        Inches(tile_top),
                        Inches(max(0.18, tile_width)),
                        Inches(tile_height),
                    )
                    tile.fill.solid()
                    tile.fill.fore_color.rgb = RGBColor(255, 255, 255)
                    tile.line.fill.background()
                    tile_tf = tile.text_frame
                    tile_tf.clear()
                    tile_p = tile_tf.paragraphs[0]
                    tile_p.alignment = PP_ALIGN.CENTER
                    tile_run = tile_p.add_run()
                    tile_run.text = detail
                    tile_run.font.size = Pt(6)
                    tile_run.font.color.rgb = RGBColor(51, 65, 85)

            ppt_shapes[key] = {
                "left": left,
                "top": top,
                "width": scaled_width,
                "height": scaled_height,
            }

        for edge in edges:
            source = ppt_shapes.get(edge.source)
            target = ppt_shapes.get(edge.target)
            if not source or not target:
                continue
            start_x = source["left"] + source["width"]
            start_y = source["top"] + (source["height"] / 2)
            end_x = target["left"]
            end_y = target["top"] + (target["height"] / 2)
            if target["left"] > source["left"] + source["width"]:
                start_x = source["left"] + source["width"]
                end_x = target["left"]
            elif source["left"] > target["left"] + target["width"]:
                start_x = source["left"]
                end_x = target["left"] + target["width"]
            elif target["top"] > source["top"]:
                start_x = source["left"] + (source["width"] / 2)
                start_y = source["top"] + source["height"]
                end_x = target["left"] + (target["width"] / 2)
                end_y = target["top"]
            else:
                start_x = source["left"] + (source["width"] / 2)
                start_y = source["top"]
                end_x = target["left"] + (target["width"] / 2)
                end_y = target["top"] + target["height"]
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.ELBOW,
                int(Inches(start_x)),
                int(Inches(start_y)),
                int(Inches(end_x)),
                int(Inches(end_y)),
            )
            connector.line.color.rgb = RGBColor(91, 100, 120)
            connector.line.width = Pt(1.0)
            if edge.style == "dashed":
                connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if edge.label:
                label_left = min(start_x, end_x) + (abs(end_x - start_x) / 2) - 0.45
                label_top = min(start_y, end_y) + (abs(end_y - start_y) / 2) - 0.08
                label_chip = slide.shapes.add_shape(
                    MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                    Inches(label_left),
                    Inches(label_top),
                    Inches(0.9),
                    Inches(0.18),
                )
                label_chip.fill.solid()
                label_chip.fill.fore_color.rgb = RGBColor(255, 255, 255)
                label_chip.line.fill.background()
                label_box = slide.shapes.add_textbox(Inches(label_left), Inches(label_top + 0.005), Inches(0.9), Inches(0.16))
                label_tf = label_box.text_frame
                label_tf.clear()
                label_p = label_tf.paragraphs[0]
                label_p.alignment = PP_ALIGN.CENTER
                label_run = label_p.add_run()
                label_run.text = edge.label
                label_run.font.size = Pt(6.5)
                label_run.font.color.rgb = RGBColor(71, 85, 105)

        buffer = BytesIO()
        prs.save(buffer)
        return base64.b64encode(buffer.getvalue()).decode("ascii")

    def _build_pptx_preview_svg(self, nodes: list[DiagramNode], edges: list[DiagramEdge]) -> str:
        width = 1600
        height = 900
        page_width = 2700
        page_height = 1450
        node_map = {node.key: node for node in nodes}
        display_keys = [node.key for node in self._ordered_nodes(nodes)]
        geometry_map = self._layout_node_geometry(nodes)
        scale_x = 1450 / page_width
        scale_y = 700 / page_height
        x_offset = 70
        y_offset = 150
        svg_parts = [
            f'<svg viewBox="0 0 {width} {height}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Generated PowerPoint slide preview">',
            f'<rect x="0" y="0" width="{width}" height="{height}" rx="22" fill="#f8fbff" stroke="#d7def3" stroke-width="2"></rect>',
            '<text x="42" y="44" font-size="28" font-weight="700" font-family="Segoe UI, Arial, sans-serif" fill="#0f172a">Enterprise Architecture Overview</text>',
            '<text x="42" y="70" font-size="13" font-family="Segoe UI, Arial, sans-serif" fill="#64748b">Native PowerPoint architecture blueprint aligned to the draw.io solution layout</text>',
            '<rect x="42" y="92" width="1510" height="6" rx="3" fill="#335cff"></rect>',
            '<rect x="52" y="132" width="1490" height="730" rx="24" fill="#ffffff" stroke="#e2e8f0" stroke-width="2"></rect>',
        ]

        for heading, x, y in self._drawio_section_headings():
            svg_parts.append(
                f'<text x="{x_offset + (x * scale_x)}" y="{y_offset + (y * scale_y) - 10}" font-size="10" font-weight="700" font-family="Segoe UI, Arial, sans-serif" fill="#475569">{escape(heading)}</text>'
            )

        positions: dict[str, tuple[float, float, float, float]] = {}
        for key in display_keys:
            node = node_map[key]
            x, y, w, h = geometry_map[key]
            sx = x_offset + (x * scale_x)
            sy = y_offset + (y * scale_y)
            sw = max(36, w * scale_x)
            sh = max(30, h * scale_y)
            positions[key] = (sx, sy, sw, sh)
            fill = f'#{self.PPTX_NODE_FILL[node.category]}'
            stroke = f'#{self.PPTX_NODE_LINE[node.category]}'
            svg_parts.append(f'<rect x="{sx}" y="{sy}" width="{sw}" height="{sh}" rx="16" fill="{fill}" stroke="{stroke}" stroke-width="2"></rect>')
            svg_parts.append(
                f'<text x="{sx + sw / 2}" y="{sy + 20}" text-anchor="middle" font-size="10.5" font-weight="700" font-family="Segoe UI, Arial, sans-serif" fill="#0f172a">{escape(node.label)}</text>'
            )
            details = self._node_details(key, node.label)[:4]
            if details:
                columns = 2 if sw > 120 else 1
                tile_gap = 6
                tile_width = (sw - 16 - ((columns - 1) * tile_gap)) / columns
                for index, detail in enumerate(details):
                    col = index % columns
                    row = index // columns
                    tile_x = sx + 8 + (col * (tile_width + tile_gap))
                    tile_y = sy + 28 + (row * 24)
                    if tile_y + 18 > sy + sh - 4:
                        break
                    svg_parts.append(f'<rect x="{tile_x}" y="{tile_y}" width="{tile_width}" height="18" rx="8" fill="#ffffff" opacity="0.94"></rect>')
                    svg_parts.append(
                        f'<text x="{tile_x + tile_width / 2}" y="{tile_y + 12}" text-anchor="middle" font-size="6.8" font-family="Segoe UI, Arial, sans-serif" fill="#334155">{escape(detail)}</text>'
                    )

        for edge in edges:
            source = positions.get(edge.source)
            target = positions.get(edge.target)
            if not source or not target:
                continue
            start_x = source[0] + source[2]
            start_y = source[1] + (source[3] / 2)
            end_x = target[0]
            end_y = target[1] + (target[3] / 2)
            mid_x = (start_x + end_x) / 2
            dash = ' stroke-dasharray="8 7"' if edge.style == "dashed" else ""
            svg_parts.append(
                f'<path d="M {start_x} {start_y} L {mid_x} {start_y} L {mid_x} {end_y} L {end_x} {end_y}" fill="none" stroke="#5b6478" stroke-width="2"{dash}></path>'
            )
            if edge.label:
                label_x = mid_x - 42
                label_y = ((start_y + end_y) / 2) - 10
                svg_parts.append(f'<rect x="{label_x}" y="{label_y}" width="84" height="18" rx="8" fill="rgba(255,255,255,0.96)"></rect>')
                svg_parts.append(
                    f'<text x="{mid_x}" y="{label_y + 12}" text-anchor="middle" font-size="7" font-family="Segoe UI, Arial, sans-serif" fill="#475569">{escape(edge.label)}</text>'
                )

        svg_parts.append("</svg>")
        return "".join(svg_parts)

    def _pptx_category_groups(self, nodes: list[DiagramNode]) -> list[dict[str, object]]:
        groups: list[dict[str, object]] = []
        for category in self.CATEGORY_ORDER:
            category_nodes = [node for node in nodes if node.category == category]
            if not category_nodes:
                continue
            groups.append(
                {
                    "category": category,
                    "label": self.CATEGORY_LABELS[category],
                    "nodes": category_nodes[:5],
                }
            )
        return groups

    def _pptx_flow_ribbons(self, edges: list[DiagramEdge]) -> list[str]:
        ribbons: list[str] = []
        seen: set[str] = set()
        for edge in edges:
            cleaned = edge.label.strip()
            if not cleaned:
                continue
            key = cleaned.casefold()
            if key in seen:
                continue
            seen.add(key)
            ribbons.append(cleaned)
        return ribbons[:5]
